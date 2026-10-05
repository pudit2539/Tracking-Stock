import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'TH Sarabun New'
        if level == 1:
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
        elif level == 2:
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Blue 900
        elif level == 3:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(5, 150, 105) # Emerald 600
    return h

def add_paragraph_styled(doc, text="", bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'TH Sarabun New'
        r_pre.font.size = Pt(14)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    if text:
        r = p.add_run(text)
        r.font.name = 'TH Sarabun New'
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet_styled(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'TH Sarabun New'
        r_pre.font.size = Pt(14)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    if text:
        r = p.add_run(text)
        r.font.name = 'TH Sarabun New'
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout_box(doc, title, text, bg_hex="F1F5F9", border_color="0284C7"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(title + "\n")
    r1.font.name = 'TH Sarabun New'
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(14, 116, 144)
    
    r2 = p.add_run(text)
    r2.font.name = 'TH Sarabun New'
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(tbl, col_widths, header_bg="1E293B"):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(tbl.columns):
        w = Inches(col_widths[i])
        for cell in col.cells:
            cell.width = w
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            
    # Header row
    for cell in tbl.rows[0].cells:
        set_cell_background(cell, header_bg)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = 'TH Sarabun New'
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    # Data rows
    for row_idx, row in enumerate(tbl.rows[1:]):
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg_color)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = 'TH Sarabun New'
                    r.font.size = Pt(13)

def build_document():
    doc = Document()
    
    # Page Setup A4
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header / Banner Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("คู่มือการใช้งานระบบติดตามสต็อกและวันหมดอายุ")
    r_title.font.name = 'TH Sarabun New'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Stock & Expiry Date Tracking System (Dairy Queen Manual)\nคู่มือฉบับสมบูรณ์: แยกบทบาท User (พนักงาน) & Admin (ผู้จัดการ) | ใช้งานบนคอมพิวเตอร์และมือถือ (LINE)")
    r_sub.font.name = 'TH Sarabun New'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # --- SECTION 1: ROLES & PERMISSIONS ---
    add_heading_styled(doc, "1. ภาพรวมระบบและสิทธิ์การใช้งาน (Roles & Permissions)", level=1)
    add_paragraph_styled(doc, "ระบบ Tracking Stock & Expiry Date ถูกออกแบบมาเพื่อร้าน Dairy Queen โดยเฉพาะ เพื่อควบคุมสต็อกวัตถุดิบและป้องกันสินค้าหมดอายุ (Zero Waste) โดยระบบเชื่อมโยงข้อมูลแบบ Real-time ระหว่างแอปพลิเคชัน LINE บนมือถือของพนักงาน และระบบ Web Dashboard บนคอมพิวเตอร์หลังร้าน")
    
    # Table of Roles
    tbl_roles = doc.add_table(rows=12, cols=3)
    role_headers = ["ฟังก์ชันการทำงาน", "พนักงานหน้าร้าน (User)", "ผู้จัดการ / แอดมิน (Admin)"]
    for i, h in enumerate(role_headers):
        tbl_roles.cell(0, i).paragraphs[0].add_run(h)
        
    role_data = [
        ("ตัดสต็อก / เบิกใช้สินค้าหน้าร้าน (FEFO)", "✅ ทำได้ (LINE & Web)", "✅ ทำได้"),
        ("รับของเข้าคลัง พร้อมระบุวันหมดอายุ", "✅ ทำได้ (LINE & Web)", "✅ ทำได้"),
        ("บันทึกของเสีย / ทิ้ง / สินค้าชำรุด (Waste)", "✅ ทำได้ (LINE & Web)", "✅ ทำได้"),
        ("เช็คสต็อกคงเหลือ / สินค้าใกล้หมด / ใกล้หมดอายุ", "✅ ทำได้ (LINE & Web)", "✅ ทำได้"),
        ("ยกเลิกคำสั่งล่าสุด (Undo ภายใน 5 นาที)", "✅ ทำได้ (พิมพ์ผ่าน LINE)", "✅ ทำได้"),
        ("ปรับ/ต่อวันหมดอายุล็อตผ่าน LINE", "✅ ทำได้ (พิมพ์ผ่าน LINE)", "✅ ทำได้"),
        ("ดูประวัติการเบิกใช้และรับเข้า (Logs)", "✅ ดูได้", "✅ ดูได้"),
        ("เพิ่มสินค้าใหม่เข้าสู่ระบบ", "❌ ไม่มีสิทธิ์", "✅ ทำได้ (ต้องใส่ PIN)"),
        ("แก้ไขข้อมูลสินค้า / ลบสินค้า", "❌ ไม่มีสิทธิ์", "✅ ทำได้ (ต้องใส่ PIN)"),
        ("จัดการทุกล็อตสินค้า (Batches Management)", "❌ ไม่มีสิทธิ์", "✅ ทำได้ (ต้องใส่ PIN)"),
        ("ตั้งค่า LINE Token / เวลาแจ้งเตือนประจำวัน", "❌ ไม่มีสิทธิ์", "✅ ทำได้ (ต้องใส่ PIN)")
    ]
    for row_idx, data in enumerate(role_data, start=1):
        for col_idx, text in enumerate(data):
            p = tbl_roles.cell(row_idx, col_idx).paragraphs[0]
            if col_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(text)
            
    format_table(tbl_roles, [3.2, 1.7, 1.8], header_bg="0F172A")
    add_paragraph_styled(doc, "")

    # --- SECTION 2: USER MANUAL ---
    add_heading_styled(doc, "2. คู่มือสำหรับพนักงานหน้าร้าน (User Manual)", level=1)
    
    add_heading_styled(doc, "2.1 การใช้งานผ่านมือถือ (LINE Bot & Rich Menu)", level=2)
    add_callout_box(doc, "📢 กฎสำคัญเวลาพิมพ์ใน LINE กลุ่ม:", 
                    "เมื่อสั่งงานใน LINE กลุ่ม ให้พิมพ์คำว่า \"DQ\" นำหน้าทุกครั้งเสมอ (เช่น DQ ตัด coke 2, DQ สั่งของ)\nเพื่อป้องกันไม่ให้บอทตอบแทรกเวลาคุยงานทั่วไป (หากพิมพ์ในแชทส่วนตัวกับบอท ไม่ต้องใส่ DQ ก็ได้)")
    
    add_heading_styled(doc, "🔘 การใช้งานปุ่มลัด Rich Menu (6 เมนูด้านล่างจอ LINE)", level=3)
    add_paragraph_styled(doc, "เมื่อเปิดแชทกับบอท พนักงานสามารถแตะ 6 ปุ่มลัดขนาดใหญ่ใต้จอได้ทันทีโดยไม่ต้องพิมพ์:")
    add_bullet_styled(doc, " ดูรายการวัตถุดิบที่เหลือต่ำกว่าเกณฑ์ Safety Stock ต้องรีบสั่งซื้อด่วน", "1. 🛒 สรุปสั่งของ:")
    add_bullet_styled(doc, " ดูรายการวัตถุดิบทุกล็อตที่จะหมดอายุใน 7 วัน พร้อมปุ่มลัดจัดการ", "2. ⏳ ใกล้หมดอายุ:")
    add_bullet_styled(doc, " ดูยอดสต็อกคงเหลือของสินค้าทุกรายการในร้าน", "3. 📊 ภาพรวมร้าน:")
    add_bullet_styled(doc, " เรียกดูตัวอย่างคำสั่งตัดสต็อกและการเบิกใช้", "4. ✂️ ตัดสต็อก / เบิก:")
    add_bullet_styled(doc, " เรียกดูตัวอย่างคำสั่งรับของเข้าคลังพร้อมระบุวันหมดอายุ", "5. 📦 รับของเข้า:")
    add_bullet_styled(doc, " กู้คืนสต็อกเดิมทันทีหากพิมพ์ตัวเลขผิด (ทำได้ภายใน 5 นาที)", "6. ↩️ ยกเลิก (Undo):")
    
    # Include Rich Menu Image if exists
    rm_path = os.path.abspath("public/img/richmenu.png")
    if os.path.exists(rm_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(rm_path, width=Inches(5.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("รูปที่ 1: แผงเมนูลัด 6 ช่อง (Rich Menu) ใต้หน้าจอแชท LINE")
        r_cap.font.name = 'TH Sarabun New'
        r_cap.font.size = Pt(11)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(100, 116, 139)

    add_heading_styled(doc, "📝 รูปแบบคำสั่งที่ใช้บ่อยในชีวิตประจำวัน", level=3)
    
    add_paragraph_styled(doc, "1) การตัดสต็อก / เบิกใช้ (FEFO - First Expired, First Out):", bold_prefix="• ")
    add_paragraph_styled(doc, "ระบบจะตัดออกจากล็อตที่หมดอายุเร็วที่สุดก่อนให้อัตโนมัติ เพื่อป้องกันสินค้าเน่าเสียคาตู้:")
    add_bullet_styled(doc, "DQ ตัด coke 2  หรือ  DQ ใช้ นมจืด 5  หรือ  DQ แก๊สบอม -1")
    
    add_paragraph_styled(doc, "2) การรับของเข้าคลัง พร้อมระบุวันหมดอายุล็อต:", bold_prefix="• ")
    add_paragraph_styled(doc, "สามารถระบุวันหมดอายุได้ทั้งแบบระบุวันที่ชัดเจน หรือบวกจำนวนเดือนจากวันนี้:")
    add_bullet_styled(doc, "DQ รับ coke 24 (รับเข้าปกติ ระบบผูกวันตามรอบสินค้า)")
    add_bullet_styled(doc, "DQ รับ แก๊สบอม 10 exp 31/12/2026 (ระบุวันหมดอายุ 31 ธ.ค. 2026)")
    add_bullet_styled(doc, "DQ รับ โคน 5 +3 เดือน (บวกวันหมดอายุเพิ่ม 3 เดือนจากวันที่รับเข้า)")

    add_paragraph_styled(doc, "3) การรับของเข้าพร้อมกันหลายรายการ (Multi-item Receive / Bulk Action):", bold_prefix="• ")
    add_paragraph_styled(doc, "สามารถสั่งรับของเข้าหลายรายการในประโยคเดียว พร้อมใส่วันหมดอายุต่อท้ายเพื่อแชร์ให้ทุกตัว:")
    add_bullet_styled(doc, "DQ รับ แก๊สบอม 10 และ โคน 10 exp 31/12/2026 (ทั้งแก๊สบอมและโคนจะได้วัน 31/12/2026)")
    add_bullet_styled(doc, "DQ รับ แก๊สบอม 10 exp 31/12/2026 และ โคน 10 exp 31/01/2027 (ระบุวันแยกแต่ละตัว)")
    add_bullet_styled(doc, "DQ รับเข้า 10 แก๊สบอม โคน exp 31/12/2026 (แบบระบุจำนวนนำหน้า)")
    add_bullet_styled(doc, "DQ รับ แก๊สบอม โคน อย่างละ 10 exp 31/12/2026 (แบบใช้คำว่าอย่างละ)")

    add_paragraph_styled(doc, "4) การบันทึกของเสีย / ทิ้ง / สินค้าชำรุด (Waste):", bold_prefix="• ")
    add_paragraph_styled(doc, "หักสต็อกและแยกประเภทลงรายงาน Waste ทันที เพื่อไม่ให้ยอดของเสียปนกับยอดขาย:")
    add_bullet_styled(doc, "DQ ทิ้ง นมจืด 1  หรือ  DQ เสีย โคน 2")

    add_paragraph_styled(doc, "5) การปรับ / เลื่อนวันหมดอายุของล็อตเดิม:", bold_prefix="• ")
    add_bullet_styled(doc, "DQ ปรับวันหมดอายุ แก๊สบอม 31/12/2026  หรือ  DQ ต่ออายุ แก๊สบอม +1 เดือน")

    add_paragraph_styled(doc, "6) การยกเลิกรายการเมื่อพิมพ์ผิด (Undo):", bold_prefix="• ")
    add_bullet_styled(doc, "พิมพ์ DQ ยกเลิก หรือ DQ undo (สามารถกู้คืนสต็อกและลบประวัติที่บันทึกผิดได้ภายใน 5 นาที)")

    add_heading_styled(doc, "2.2 การใช้งานผ่านคอมพิวเตอร์ (Web Browser สำหรับพนักงาน)", level=2)
    add_paragraph_styled(doc, "สำหรับพนักงานที่อยู่หน้าเคาน์เตอร์แคชเชียร์หรือคอมพิวเตอร์ร้าน:")
    add_bullet_styled(doc, " เปิด Google Chrome หรือ Browser แล้วไปที่ URL ระบบ (เช่น https://tracking-stock.vercel.app)", "1. เข้าสู่ระบบ:")
    add_bullet_styled(doc, " ตรวจดูตัวเลขสต็อกคงเหลือ แถบสีเตือน (เขียว = ปกติ, ส้ม = ต่ำกว่า Safety Stock, แดง = หมด) และวันหมดอายุของล็อตที่ใกล้หมดที่สุด", "2. ค้นหาและดูคลังสินค้า:")
    add_bullet_styled(doc, " กดปุ่มเครื่องหมาย \"+\" (รับเข้า) หรือ \"-\" (เบิกใช้) ที่แถวสินค้านั้น ใส่จำนวนแล้วกดยืนยัน", "3. ตัดสต็อก / รับเข้าด่วน:")
    add_bullet_styled(doc, " สลับไปแท็บ \"ประวัติการเบิกใช้\" หรือ \"ประวัติการรับเข้า\" เพื่อดูย้อนหลังได้ทันที", "4. ตรวจสอบประวัติ:")

    # --- SECTION 3: ADMIN MANUAL ---
    add_heading_styled(doc, "3. คู่มือสำหรับผู้จัดการและผู้ดูแลระบบ (Admin Manual)", level=1)
    
    add_heading_styled(doc, "3.1 การเข้าสู่โหมด Admin บนคอมพิวเตอร์", level=2)
    add_paragraph_styled(doc, "เพื่อความปลอดภัย หน้าเว็บจะเริ่มต้นด้วย \"โหมดพนักงาน (Staff Mode)\" ซึ่งจะซ่อนปุ่มแก้ไข ลบ หรือตั้งค่าสำคัญไว้:")
    add_bullet_styled(doc, " ให้คลิกที่ปุ่มสีเทาเขียนว่า \"โหมดพนักงาน (Staff)\" ที่มุมบนขวาของหน้าจอ", "ขั้นตอนการปลดล็อก:")
    add_bullet_styled(doc, " หน้าต่างใส่รหัส PIN จะปรากฏขึ้น ให้กรอกรหัส Admin PIN (ค่าเริ่มต้น: 1234 หรือตามที่ผู้จัดการตั้งไว้)")
    add_bullet_styled(doc, " เมื่อรหัสถูกต้อง แถบจะเปลี่ยนเป็นสีเขียว \"โหมดผู้จัดการ (Admin Active)\" พร้อมปลดล็อกปุ่มจัดการทั้งหมดทันที")

    add_heading_styled(doc, "3.2 การสร้างสินค้าใหม่ พร้อมใส่สต็อกเริ่มต้นและวันหมดอายุ", level=2)
    add_paragraph_styled(doc, "เมื่อมีสินค้าหรือวัตถุดิบตัวใหม่เข้ามาในสาขา:")
    add_bullet_styled(doc, " เข้าสู่โหมด Admin แล้วคลิกปุ่ม \"+ เพิ่มสินค้าใหม่\" (ปุ่มสีน้ำเงินมุมขวาบน)", "ขั้นตอนที่ 1:")
    add_bullet_styled(doc, " กรอกชื่อสินค้า, เลือกหมวดหมู่, ใส่หน่วยนับ (เช่น กล่อง, ถุง, ลัง, ชิ้น)", "ขั้นตอนที่ 2:")
    add_bullet_styled(doc, " กำหนดเกณฑ์สั่งซื้อขั้นต่ำ (Safety Stock) และจำนวนวันเตือนหมดอายุล่วงหน้า (เช่น 7 วัน)", "ขั้นตอนที่ 3:")
    add_bullet_styled(doc, " ใส่วันหมดอายุของล็อตแรก และจำนวนสต็อกเริ่มต้น (Initial Stock) ได้ในขั้นตอนนี้ทันที โดยไม่ต้องไปกดรับเข้าแยกอีกครั้ง", "ขั้นตอนที่ 4:")
    add_bullet_styled(doc, " กดปุ่ม \"บันทึกสินค้า\" ➔ สินค้าและล็อตแรกจะถูกบันทึกเข้าคลังพร้อมทำงานทันที", "ขั้นตอนที่ 5:")

    add_heading_styled(doc, "3.3 การจัดการทุกล็อตสินค้าและวันหมดอายุ (Batches Management)", level=2)
    add_callout_box(doc, "🛡️ ข้อควรรู้เรื่องการแยกล็อต (Batch Isolation 100%):", 
                    "ระบบเก็บข้อมูลแบบแยกล็อตเด็ดขาด หากรับเข้าสินค้าตัวเดียวกันแต่วันหมดอายุต่างกัน (เช่น ล็อต 2026 และ ล็อต 2027)\nระบบจะแยกเป็นคนละแถว มีเลขล็อต จำนวน และวันหมดอายุของตัวเอง ไม่มีการนำมาทับหรือปนกันแน่นอน\nและเวลาตัดสต็อก ระบบจะตัดจากล็อต 2026 จนหมดเกลี้ยงก่อน แล้วค่อยขยับไปตัดล็อต 2027 อัตโนมัติ (FEFO)")
    
    add_paragraph_styled(doc, "วิธีตรวจสอบและแก้ไขล็อตรายตัวบนหน้าเว็บ:")
    add_bullet_styled(doc, " ในหน้ารายการสินค้า ให้คลิกที่ไอคอนกล่อง \"จัดการล็อต (Batches)\" ที่แถวสินค้านั้น")
    add_bullet_styled(doc, " หน้าต่างจะแสดงทุกล็อตที่ยังมีของอยู่ พร้อมเลขล็อต, วันที่รับเข้า, วันหมดอายุ, และป้ายสถานะสี")
    add_bullet_styled(doc, " สามารถแก้ไขตัวเลขสต็อก หรือแก้วันหมดอายุเฉพาะล็อตนั้นได้โดยตรง แล้วกดยืนยัน")
    add_bullet_styled(doc, " หากมีล็อตผิดพลาด สามารถกดไอคอนถังขยะเพื่อลบล็อตนั้นทิ้งได้ทันที")

    add_heading_styled(doc, "3.4 การจัดการพร้อมกันหลายรายการ (Bulk Actions)", level=2)
    add_paragraph_styled(doc, "เหมาะสำหรับการตรวจนับสต็อกสิ้นเดือน หรือการรับสินค้าลงคลังพร้อมกันหลายสิบรายการ:")
    add_bullet_styled(doc, " ติ๊กเครื่องหมายถูกหน้ารายการสินค้าที่ต้องการ (หรือกดติ๊กที่หัวตารางเพื่อเลือกทั้งหมด)", "ขั้นตอนที่ 1:")
    add_bullet_styled(doc, " แถบเครื่องมือสีเข้ม Bulk Actions จะลอยขึ้นมาด้านล่างจอ:", "ขั้นตอนที่ 2:")
    add_bullet_styled(doc, "มีปุ่มด่วน +1, +5, +10, +20, +50 สามารถกดคลิกซ้ำๆ เพื่อบวกยอดสะสมได้เรื่อยๆ (เช่น คลิก +10 สามครั้ง ยอดกลายเป็น +30)", "   • รับเข้ากลุ่ม (Add Stock):")
    add_bullet_styled(doc, "ใส่วันที่ หรือกดปุ่มลัด +1 ด., +3 ด., +6 ด., +1 ปี เพื่อตั้งวันหมดอายุให้ทุกล็อตที่เลือกพร้อมกันในคลิกเดียว", "   • กำหนดวันหมดอายุกลุ่ม:")
    add_bullet_styled(doc, "ตั้งค่า Safety Stock ขั้นต่ำเท่ากันสำหรับกลุ่มสินค้านั้น", "   • ปรับเกณฑ์ Safety Stock:")
    add_bullet_styled(doc, "ลบสินค้าที่ไม่ใช้งานพร้อมกันทีละหลายรายการอย่างปลอดภัย", "   • ลบสินค้ากลุ่ม:")
    add_bullet_styled(doc, " กดปุ่ม \"บันทึกการเปลี่ยนแปลงทั้งหมด\"", "ขั้นตอนที่ 3:")

    add_heading_styled(doc, "3.5 การตั้งค่าระบบและการแจ้งเตือน LINE", level=2)
    add_paragraph_styled(doc, "เข้าที่แท็บ \"⚙️ ตั้งค่าระบบ (Settings)\":")
    add_bullet_styled(doc, " วาง Channel Access Token และ Target ID (Group ID หรือ User ID) แล้วกดปุ่ม \"🔔 ทดสอบส่งข้อความ\"", "1. การเชื่อมต่อ LINE:")
    add_bullet_styled(doc, " กำหนดเวลาที่ต้องการให้บอทสรุปยอดส่งเข้า LINE อัตโนมัติ (เช่น 08:00 น. ของทุกเช้า)", "2. เวลาแจ้งเตือนประจำวัน:")
    add_bullet_styled(doc, " คลิกปุ่ม \"🎨 ติดตั้ง Rich Menu บน LINE\" ระบบจะสร้างและติดตั้งเมนู 6 ช่องที่สมบูรณ์ให้ทันที", "3. ติดตั้ง Rich Menu:")
    add_bullet_styled(doc, " สามารถเปลี่ยนรหัสผ่านโหมด Admin ได้ในหน้านี้", "4. เปลี่ยนรหัส Admin PIN:")

    # Include Line Preview Image if exists
    lp_path = os.path.abspath("public/line_preview.png")
    if os.path.exists(lp_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(6)
        p_img2.paragraph_format.space_after = Pt(2)
        doc.add_picture(lp_path, width=Inches(3.8))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap2 = p_cap2.add_run("รูปที่ 2: ตัวอย่างการ์ด Flex Message แจ้งเตือนสรุปสต็อกและวันหมดอายุใน LINE")
        r_cap2.font.name = 'TH Sarabun New'
        r_cap2.font.size = Pt(11)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = RGBColor(100, 116, 139)

    add_heading_styled(doc, "3.6 การส่งออกรายงาน Excel (UTF-8 BOM CSV ภาษาไทย 100%)", level=2)
    add_paragraph_styled(doc, "ในแท็บรายงาน ผู้จัดการสามารถกดปุ่ม \"📥 ส่งออก Excel (CSV)\" เพื่อดาวน์โหลดข้อมูล:")
    add_bullet_styled(doc, " รหัสสินค้า, ชื่อสินค้า, หมวดหมู่, สต็อกคงเหลือ, สถานะ, และวันหมดอายุใกล้สุด", "1. ข้อมูลสต็อกคงเหลือปัจจุบัน:")
    add_bullet_styled(doc, " วันที่รับ, เลขล็อต, ผู้รับเข้า, จำนวน, และวันหมดอายุ", "2. ประวัติการรับของเข้า (Inbound Logs):")
    add_bullet_styled(doc, " วันที่เบิก, เลขล็อตที่ถูกตัด, ผู้เบิกใช้, สาเหตุ (เบิกใช้ หรือ ของเสีย)", "3. ประวัติการตัดสต็อกและของเสีย (Usage & Waste Logs):")
    add_bullet_styled(doc, " รายการสินค้าที่สต็อกต่ำกว่าเกณฑ์ พร้อมจำนวนที่ระบบแนะนำให้สั่งซื้อ", "4. สรุปรายการของต้องสั่งซื้อ (Purchase Order List):")
    add_paragraph_styled(doc, "💡 ไฟล์รายงานถูกเข้ารหัสแบบ UTF-8 BOM ทำให้เปิดในโปรแกรม Microsoft Excel บน Windows และ Mac ได้ทันที ภาษาไทยสมบูรณ์ สระไม่ลอย ไม่เป็นภาษาต่างดาว")

    # --- SECTION 4: CHEATSHEET ---
    doc.add_page_break()
    add_heading_styled(doc, "4. ตารางสรุปคำสั่งลัด LINE บอท (Quick Cheatsheet)", level=1)
    add_paragraph_styled(doc, "🖨️ สามารถพิมพ์หน้านี้ไปติดไว้ที่ผนังครัวหรือเคาน์เตอร์หน้าร้าน เพื่อให้พนักงานดูเป็นแนวทางในการสั่งงานบอทได้ทันที")
    
    tbl_cheat = doc.add_table(rows=14, cols=3)
    cheat_headers = ["หมวดหมู่คำสั่ง", "ตัวอย่างคำสั่งพิมพ์ใน LINE (ในกลุ่มใส่ DQ)", "ผลลัพธ์การทำงานของระบบ"]
    for i, h in enumerate(cheat_headers):
        tbl_cheat.cell(0, i).paragraphs[0].add_run(h)
        
    cheat_data = [
        ("✂️ ตัดสต็อก / เบิกใช้", "DQ ตัด coke 2\nDQ ใช้ นมจืด 5\nDQ โอริโอ้ -3", "ตัดสต็อกจากล็อตที่หมดอายุก่อนให้อัตโนมัติ (FEFO)"),
        ("📦 รับของเข้า (ระบุวัน)", "DQ รับ แก๊สบอม 10 exp 31/12/2026\nDQ รับ โคน 5 +3 เดือน\nDQ รับ coke 24", "เพิ่มสต็อกและสร้างล็อตใหม่พร้อมวันหมดอายุที่ระบุ"),
        ("⚡ รับหลายอย่างพร้อมกัน", "DQ รับ แก๊สบอม 10 และ โคน 10 exp 31/12/2026\nDQ รับเข้า 10 แก๊สบอม โคน exp 31/12/2026\nDQ รับ แก๊สบอม โคน อย่างละ 10 exp 31/12/2026", "รับเข้าพร้อมกันหลายตัว และผูกวันหมดอายุให้ทุกตัวพร้อมกัน"),
        ("🗑️ บันทึกของเสีย (ทิ้ง)", "DQ ทิ้ง นมสด 1\nDQ เสีย โคน 2\nDQ waste ช้อน 5", "หักสต็อกและแยกประเภทเป็น \"ของเสีย\" เพื่อดูยอด Loss"),
        ("📅 ปรับวันหมดอายุ", "DQ ปรับวันหมดอายุ แก๊สบอม 31/12/2026\nDQ ต่ออายุ แก๊สบอม +1 เดือน", "แก้ไขวันหมดอายุของล็อตสินค้า"),
        ("📊 นับสต็อกจริง", "DQ coke เหลือ 10\nDQ นับ นมจืด ได้ 8", "ปรับยอดสต็อกคงเหลือรวมให้ตรงกับยอดจริงที่นับได้"),
        ("🔍 เช็คยอดสินค้า", "DQ เช็ค coke\nDQ นมจืด เหลือเท่าไหร่", "ส่งการ์ดข้อมูลสต็อก วันหมดอายุ และลิงก์ดูบนเว็บ"),
        ("🛒 เช็คของต้องสั่ง", "DQ สั่งของ  หรือ  DQ ของหมด", "สรุปสินค้าที่ต่ำกว่า Safety Stock พร้อมจำนวนที่ควรสั่ง"),
        ("⏳ เช็คของใกล้หมดอายุ", "DQ ใกล้หมดอายุ  หรือ  DQ หมดอายุ", "ส่งการ์ดเตือนสินค้าทุกล็อตที่กำลังจะหมดอายุใน 7 วัน"),
        ("📈 สรุปภาพรวมร้าน", "DQ ภาพรวม  หรือ  DQ สรุป", "ส่งการ์ดสรุปสถานะทั้งร้าน (ปกติ, ใกล้หมด, ใกล้หมดอายุ)"),
        ("↩️ ยกเลิกคำสั่ง (Undo)", "DQ ยกเลิก  หรือ  DQ undo", "กู้คืนสต็อกเดิมทันทีหากเพิ่งพิมพ์ผิด (ภายใน 5 นาที)"),
        ("🔘 เมนูลัด Rich Menu", "แตะ 6 ปุ่มลัดใต้หน้าจอแชท", "กดสั่งงานลัดใน 1 วินาที ไม่ต้องพิมพ์"),
        ("❓ ดูคู่มือช่วยเหลือ", "DQ วิธีใช้  หรือ  DQ", "บอทส่งข้อความแนะนำคำสั่งลัดและปุ่มตัวช่วย")
    ]
    
    for row_idx, data in enumerate(cheat_data, start=1):
        for col_idx, text in enumerate(data):
            p = tbl_cheat.cell(row_idx, col_idx).paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(text)
            
    format_table(tbl_cheat, [1.8, 2.7, 2.2], header_bg="047857") # Emerald 700
    
    output_filename = "คู่มือการใช้งาน_ระบบTracking_Stock_DairyQueen.docx"
    doc.save(output_filename)
    print("Document created successfully in project root.")
    
    # Also save a copy to the artifacts directory
    artifact_dir = r"C:\Users\pudit.ye\.gemini\antigravity\brain\7b31d491-ebba-4dfc-8eb1-dcfdaf0c47ee"
    if os.path.exists(artifact_dir):
        artifact_path = os.path.join(artifact_dir, output_filename)
        doc.save(artifact_path)
        print("Document copy saved to artifacts successfully.")

if __name__ == "__main__":
    build_document()
